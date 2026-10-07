import csv
import random
import numpy as np
import openpyxl

# Set seeds for reproducibility
np.random.seed(42)
random.seed(42)

def generate_placement_dataset(n_samples=2500):
    departments = ['Computer Science', 'Information Technology', 'Electronics & Comm', 
                   'Electrical & Electronics', 'Mechanical', 'Civil']
    dept_weights = [0.35, 0.25, 0.15, 0.10, 0.10, 0.05]
    
    first_names = ['Aarav', 'Ananya', 'Rohan', 'Priya', 'Vikram', 'Neha', 'Aditya', 'Sneha', 
                   'Rahul', 'Kavya', 'Siddharth', 'Pooja', 'Arjun', 'Ishita', 'Dev', 'Riya',
                   'Karan', 'Meera', 'Varun', 'Divya', 'Amit', 'Tanvi', 'Yash', 'Shruti']
    last_names = ['Sharma', 'Verma', 'Patel', 'Rao', 'Nair', 'Gupta', 'Singh', 'Reddy',
                  'Kumar', 'Joshi', 'Mehta', 'Iyengar', 'Chatterjee', 'Deshmukh', 'Chawla']
    
    names = [f"{random.choice(first_names)} {random.choice(last_names)}" for _ in range(n_samples)]
    depts = np.random.choice(departments, size=n_samples, p=dept_weights)
    
    # 1. CGPA (5.0 to 9.9)
    cgpa = np.round(np.random.beta(a=5, b=2, size=n_samples) * 5 + 5, 2)
    cgpa = np.clip(cgpa, 5.0, 9.9)
    
    # 2. LeetCode Ratings & Solved
    leetcode_solved = np.random.geometric(p=0.005, size=n_samples) + np.random.randint(0, 50, size=n_samples)
    leetcode_solved = np.clip(leetcode_solved, 0, 750)
    
    leetcode_rating = (800 + leetcode_solved * 1.8 + np.random.normal(0, 50, size=n_samples)).astype(int)
    leetcode_rating = np.clip(leetcode_rating, 800, 2400)
    
    # 3. GitHub Activity Metrics
    github_repos = np.random.poisson(lam=6, size=n_samples) + np.random.randint(0, 5, size=n_samples)
    github_repos = np.clip(github_repos, 0, 45)
    
    github_commits_year = (github_repos * 18 + np.random.exponential(scale=100, size=n_samples)).astype(int)
    github_commits_year = np.clip(github_commits_year, 0, 1200)
    
    github_stars = (github_repos * 1.5 + np.random.exponential(scale=4, size=n_samples)).astype(int)
    github_stars = np.clip(github_stars, 0, 150)
    
    github_score = np.round(
        np.clip(
            (github_repos / 30 * 30) + 
            (github_commits_year / 600 * 50) + 
            (github_stars / 40 * 20), 0, 100
        ), 1
    )
    
    # 4. LinkedIn Activity Metrics
    linkedin_connections = np.random.randint(30, 500, size=n_samples)
    linkedin_certs = np.random.poisson(lam=2, size=n_samples)
    linkedin_certs = np.clip(linkedin_certs, 0, 10)
    
    linkedin_posts_freq = np.random.randint(1, 11, size=n_samples)
    
    linkedin_score = np.round(
        np.clip(
            (linkedin_connections / 500 * 40) + 
            (linkedin_certs / 6 * 40) + 
            (linkedin_posts_freq / 10 * 20), 0, 100
        ), 1
    )
    
    # 5. Real World Projects & Skills
    projects_count = np.random.poisson(lam=3, size=n_samples)
    projects_count = np.clip(projects_count, 0, 10)
    
    complexity_options = ['Basic', 'Intermediate', 'Advanced', 'Production-Grade']
    complexity_weights = [0.25, 0.45, 0.22, 0.08]
    project_complexity = np.random.choice(complexity_options, size=n_samples, p=complexity_weights)
    
    complexity_map = {'Basic': 25, 'Intermediate': 55, 'Advanced': 80, 'Production-Grade': 100}
    proj_comp_numeric = np.array([complexity_map[c] for c in project_complexity])
    
    project_score = np.round(
        np.clip(
            (projects_count / 6 * 40) + (proj_comp_numeric * 0.6), 0, 100
        ), 1
    )
    
    skills_count = np.random.randint(2, 12, size=n_samples)
    skill_score = np.round(np.clip(skills_count * 7 + np.random.normal(25, 5, size=n_samples), 10, 100), 1)
    
    # 6. Formula for Placement Readiness
    coding_score = np.clip((leetcode_rating - 800) / 1400 * 60 + (leetcode_solved / 500 * 40), 0, 100)
    cgpa_score = (cgpa / 10.0) * 100
    
    dept_boost = np.array([5.0 if d in ['Computer Science', 'Information Technology'] else (2.0 if d == 'Electronics & Comm' else 0.0) for d in depts])
    
    raw_readiness = (
        0.28 * cgpa_score +
        0.32 * coding_score +
        0.18 * project_score +
        0.10 * github_score +
        0.07 * linkedin_score +
        0.05 * skill_score +
        dept_boost +
        np.random.normal(0, 1.2, size=n_samples)
    )
    
    readiness_pct = np.round(np.clip(raw_readiness, 5.0, 99.5), 1)
    
    placement_tier = []
    placement_status_label = []
    for score in readiness_pct:
        if score < 52.0:
            placement_tier.append(0)
            placement_status_label.append('Not Ready')
        elif score < 68.0:
            placement_tier.append(1)
            placement_status_label.append('Service Tier Ready')
        elif score < 83.0:
            placement_tier.append(2)
            placement_status_label.append('Product Tier Ready')
        else:
            placement_tier.append(3)
            placement_status_label.append('Dream Tier Ready')
            
    header = [
        'student_id', 'name', 'dept', 'cgpa', 'skills_count', 'skill_score',
        'leetcode_rating', 'leetcode_solved', 'github_repos', 'github_commits_year',
        'github_stars', 'github_score', 'linkedin_connections', 'linkedin_certs',
        'linkedin_posts_freq', 'linkedin_score', 'projects_count', 'project_complexity',
        'project_score', 'readiness_score', 'placement_tier', 'placement_tier_label'
    ]
    
    rows = []
    for i in range(n_samples):
        rows.append([
            f"STU-{1000+i}", names[i], depts[i], float(cgpa[i]), int(skills_count[i]), float(skill_score[i]),
            int(leetcode_rating[i]), int(leetcode_solved[i]), int(github_repos[i]), int(github_commits_year[i]),
            int(github_stars[i]), float(github_score[i]), int(linkedin_connections[i]), int(linkedin_certs[i]),
            int(linkedin_posts_freq[i]), float(linkedin_score[i]), int(projects_count[i]), project_complexity[i],
            float(project_score[i]), float(readiness_pct[i]), int(placement_tier[i]), placement_status_label[i]
        ])
        
    # Write CSV
    csv_path = 'placement_dataset.csv'
    with open(csv_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(rows)
    print(f"Dataset generated with {n_samples} records saved to '{csv_path}'.")
    
    # Write Excel summary via openpyxl
    excel_path = 'placement_dataset_info.xlsx'
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sample_Records"
    ws1.append(header)
    for r in rows[:100]:
        ws1.append(r)
        
    ws2 = wb.create_sheet(title="Department_Summary")
    ws2.append(["Department", "Total Students", "Avg CGPA", "Avg LeetCode Rating", "Avg Readiness Score (%)"])
    
    dept_stats = {d: {'count': 0, 'cgpa_sum': 0.0, 'lc_sum': 0.0, 'score_sum': 0.0} for d in departments}
    for r in rows:
        d = r[2]
        dept_stats[d]['count'] += 1
        dept_stats[d]['cgpa_sum'] += r[3]
        dept_stats[d]['lc_sum'] += r[6]
        dept_stats[d]['score_sum'] += r[19]
        
    for d, st in dept_stats.items():
        cnt = st['count']
        if cnt > 0:
            ws2.append([d, cnt, round(st['cgpa_sum']/cnt, 2), round(st['lc_sum']/cnt, 1), round(st['score_sum']/cnt, 1)])
            
    ws3 = wb.create_sheet(title="Placement_Tier_Summary")
    ws3.append(["Placement Tier", "Student Count"])
    tier_counts = {}
    for r in rows:
        lbl = r[21]
        tier_counts[lbl] = tier_counts.get(lbl, 0) + 1
    for lbl, count in tier_counts.items():
        ws3.append([lbl, count])
        
    wb.save(excel_path)
    print(f"Excel summary report saved to '{excel_path}'.")
    return rows, header

if __name__ == '__main__':
    generate_placement_dataset()
